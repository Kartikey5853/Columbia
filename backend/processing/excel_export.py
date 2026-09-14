from __future__ import annotations

from io import BytesIO

from .product_schema import MARKETPLACES, availability_display, format_inr, normalize_sku, price_value


SITE_LABELS = {
    "amazon": "Amazon", "ajio": "AJIO", "columbia": "Columbia",
    "adventuras": "Adventuras", "adventure": "Adventuras", "myntra": "Myntra",
    "tatacliq": "TataCliQ", "tata_lux": "Tata Lux", "flipkart": "Flipkart",
}
EXPORT_MARKETPLACES = ("amazon", "ajio", "columbia", "adventuras", "myntra", "tatacliq", "tata_lux", "flipkart")


def _value(card: dict | None, field: str):
    if not isinstance(card, dict):
        return None
    return card.get(field)


def _availability_value(row: dict, source: str, card: dict | None, value):
    # Status records are historical and may be false simply because no match
    # exists.  Only explicit source availability can produce OOS.
    return availability_display(source, card, value)


def tuple_export_rows(products: dict, options: dict[str, bool]) -> list[dict]:
    rows: list[dict] = []
    for ean, row in sorted(products.items()):
        item: dict = {}
        if options.get("identifiers", True):
            item["Canonical Product ID"] = row.get("canonical_product_id")
            item["EAN"] = row.get("EAN")
            item["Columbia SKU"] = normalize_sku(row.get("columbia_sku") or _value(row.get("columbia"), "sku"))
            item["Columbia Product ID"] = row.get("columbia_product_id") or _value(row.get("columbia"), "source_product_id")
        if options.get("source_ids", True):
            for source in EXPORT_MARKETPLACES:
                card = row.get(source) or (row.get("adventure") if source == "adventuras" else None)
                item[f"{SITE_LABELS[source]} Product ID"] = _value(card, "source_product_id")
                item[f"{SITE_LABELS[source]} SKU"] = normalize_sku(_value(card, "sku"))
        if options.get("prices", True):
            for source in EXPORT_MARKETPLACES:
                card = row.get(source) or (row.get("adventure") if source == "adventuras" else None)
                raw_price = _value(card, "normal_price") or _value(card, "price") or _value(card, "price_value")
                num_price = price_value(raw_price)
                item[f"{SITE_LABELS[source]} Price"] = _availability_value(row, source, card, num_price)
        if options.get("special_prices", True):
            card = row.get("ajio")
            offer = price_value(_value(card, "offer_price"))
            item["AJIO Special Price"] = _availability_value(row, "ajio", card, offer)
        if options.get("titles", True):
            for source in EXPORT_MARKETPLACES:
                card = row.get(source) or (row.get("adventure") if source == "adventuras" else None)
                item[f"{SITE_LABELS[source]} Title"] = _value(card, "title")
        if options.get("urls", True):
            for source in EXPORT_MARKETPLACES:
                card = row.get(source) or (row.get("adventure") if source == "adventuras" else None)
                item[f"{SITE_LABELS[source]} URL"] = _value(card, "url")
        if options.get("image_urls", True):
            for source in EXPORT_MARKETPLACES:
                card = row.get(source) or (row.get("adventure") if source == "adventuras" else None)
                item[f"{SITE_LABELS[source]} Image URL"] = _value(card, "image")
        rows.append(item)
    return rows


def excel_bytes(rows: list[dict], sheet_name: str = "tuples") -> bytes:
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name[:31]
    headers = list(rows[0]) if rows else []
    sheet.append(headers or ["No data"])
    for row in rows:
        sheet.append([row.get(header) for header in headers])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def excel_bytes_by_site(rows: list[dict]) -> bytes:
    """Export tuple rows into one worksheet per marketplace."""
    from openpyxl import Workbook

    sites = [
        ("amazon", "Amazon"), ("ajio", "AJIO"), ("columbia", "Columbia"),
        ("adventuras", "Adventuras"), ("myntra", "Myntra"), ("tatacliq", "TataCliq"),
        ("tata_lux", "Tata Lux"), ("flipkart", "Flipkart")
    ]
    workbook = Workbook()
    workbook.remove(workbook.active)
    common = ["Columbia SKU", "EAN(s)", "Product Image"]
    for key, label in sites:
        sheet = workbook.create_sheet(label[:31])
        prefix = f"{label} "
        site_headers = [header for header in (rows[0].keys() if rows else []) if header.startswith(prefix)]
        headers = common + site_headers
        sheet.append(headers)
        for row in rows:
            sheet.append([row.get(header) for header in headers])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
