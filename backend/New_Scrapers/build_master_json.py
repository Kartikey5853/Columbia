"""Convert master_excel.xlsx into a reusable identifier mapping."""
import json
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent
wb = openpyxl.load_workbook(ROOT / "master_excel.xlsx", read_only=True, data_only=True)
ws = wb.active
headers = ["" if c.value is None else str(c.value).strip() for c in next(ws.iter_rows())]
products = []
for values in ws.iter_rows(min_row=2, values_only=True):
    row = {headers[i]: "" if values[i] is None else str(values[i]).strip() for i in range(len(headers))}
    if row.get("EAN CODE") or row.get("SKU CODE"): products.append(row)
(ROOT / "master.json").write_text(json.dumps({"schema_version": 1, "source": "master_excel.xlsx", "columns": headers, "products": products}, indent=2), encoding="utf-8")
print(f"Wrote {ROOT / 'master.json'} with {len(products)} rows")
import sys
sys.path.insert(0, str(ROOT.parent))
from pipeline.master_pipeline import build_master_tuples
result = build_master_tuples()
print(f"Rebuilt tuple viewer data with {result['summary']['normalized_products']} rows")
